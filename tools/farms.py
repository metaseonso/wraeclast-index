"""Farming strategies for the Farms tab: BawLoch's public tier list sheet -> data/farms.json.

Source: "0.5 Tier List - by Bawloch", a public Google Sheet (no login needed):
  https://docs.google.com/spreadsheets/d/1slfzZiehPZLnR7ADk6LouuTMb8KRZDVTAUmWdjMc2Dk
The sheet is not a table. A "Tierlist" tab puts each strategy in a tier band and links to its own tab;
each strategy tab is a fixed layout of merged cells (notes, waystone, tablets, atlas master, difficulty,
investment). It has no profit numbers, and none are added here: the page prices each item live from
data/market.json.

Added on top of the cells, because the sheet has no field for them:
  - mechanics per strategy (MECH), read from the strategy name, its tablets and notes
  - the items each strategy spends and makes, under the names the market and the official trade site use.
    The waystone tier and the tablets (one per slot) are read from their cells; the rest is listed in ITEMS,
    and each entry's wording is checked against the tab so a changed sheet shows up as a warning.
Fixes to the sheet text, only where the intent is certain:
  - spelling (TYPOS): "Pradise" -> "Paradise", missing apostrophes in item names ("Aldurs Saga" -> "Aldur's Saga"),
    "diminshing", "alongisde", "atleast", "ontop", "a recent patched", "complimentary mods"
  - the Maxroll guide cell on "200% Doryani Deli Farming" and "Deli Splinter Farming" says "by aer0" and links to
    the Ritual guide (copied from the Ritual tab), so that link is left out there (WRONG_GUIDE)
  - a cut-off tablet line on Deli Boss Rush ("Delirium spawns % increased .") is left out; the other slots
    read "... increased Mirror shards." but the cut one is not guessed
  - "Vid by Me :)" is the sheet author's own video
  - links pasted into notes become buttons with a plain label (LINK_LABELS); the note keeps its own words

Rolled tablets and waystones cost more than plain ones, so each gets an official trade search built from the
mods the sheet asks for (STAT_RULES, stat ids from data/trade.json) and written to data/farmqueries.json as
{key, label, query}. The sheet's own trade links are saved searches for an old league that cannot be read back,
so the searches are rebuilt. The site's worker runs them a few at a time through the hour and serves the
median of the 10 cheapest listings at data/farmprices.json. Plain items keep their poe.ninja price.

Usage:  python tools/farms.py              download the sheet, write data/farms.json and data/farmqueries.json
        python tools/farms.py sheet.xlsx   read a saved xlsx export instead of downloading it
Needs openpyxl:  pip install openpyxl
"""
import collections
import datetime as dt
import html
import io
import json
import re
import sys
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SHEET_ID = '1slfzZiehPZLnR7ADk6LouuTMb8KRZDVTAUmWdjMc2Dk'
BASE = 'https://docs.google.com/spreadsheets/d/' + SHEET_ID
UA = 'wraeclast-index/1.0 (contact: https://github.com/metaseonso/wraeclast-index/issues)'
OUT = ROOT / 'data' / 'farms.json'
QUERIES = ROOT / 'data' / 'farmqueries.json'

TYPOS = [
    (r'\bPradise\b', 'Paradise', 0),
    (r'\bAldurs Saga', "Aldur's Saga", re.I),
    (r'\bRakiatas Flow', "Rakiata's Flow", re.I),
    (r'\bGarukhans Resolve', "Garukhan's Resolve", re.I),
    (r'\bdiminshing\b', 'diminishing', 0),
    (r'\balongisde\b', 'alongside', 0),
    (r'\batleast\b', 'at least', 0),
    (r'\bontop\b', 'on top', 0),
    (r'\ba recent patched\b', 'a recent patch', 0),
    (r'\bcomplimentary mods\b', 'complementary mods', 0),
]
WRONG_GUIDE = {'200% Doryani Deli Farming', 'Deli Splinter Farming'}
CUT_LINE = re.compile(r'%\s*increased\s*\.$')      # a tablet mod that stops mid-line
LINK_LABELS = {                                     # links pasted into notes: what they show
    'prnt.sc/Dw8q9SbpP5OK': 'Genesis tree for catalysts',
    'prnt.sc/Vgdj-qWUgD2c': 'Good tablet mods',
    'prnt.sc/8tx3nn6Sfn3f': 'Catalyst tree',
    'prnt.sc/b3k_3bJW8UEA': 'Amulet tree',
    'prnt.sc/_tev9eXwl7tB': 'Boss map list',
    '1Mm5X0CzXL2uBRfpBkWyN09hqrmDy5qXQkGxExzviuyY': 'Loot test sheet',
}
AUTHOR_LINKS = [('youtube.com', 'YouTube'), ('twitch.tv', 'Twitch'), ('maxroll.gg', 'Maxroll'), ('discord.gg', 'Discord')]

# Mechanics, by tab (or by name for the tier-list-only entries). The sheet has no mechanic field.
MECH = {
    'Deli Unstable Breach': ['Delirium', 'Breach'],
    'Lineage/Fragment Farming': ['Unique maps', 'Bosses'],
    'Deli Abyss': ['Delirium', 'Abyss'],
    'Jackpot Ritual': ['Ritual'],
    'Ritual': ['Ritual'],
    'Tablet Farming': ['Tablets'],
    'Aldurs Saga Hunting': ['Expedition'],
    'Big Boom Expedition': ['Expedition'],
    'Grand Expedition Farming': ['Expedition'],
    'Hilda Logbook Farming': ['Expedition'],
    'Unstable Hiveblood farming': ['Breach'],
    'Deli Boss Rush': ['Delirium', 'Bosses'],
    '200% Doryani Deli Farming': ['Delirium'],
    'Deli Splinter Farming': ['Delirium', 'Simulacrum'],
    'Liquid Emotion Farming': ['Delirium', 'Temple', 'Bosses'],
    'Abyssal Boxes': ['Abyss'],
    'Simulacrum': ['Delirium', 'Simulacrum'],
    'Atziri Boss Rush': ['Temple', 'Bosses'],
    'Essences': ['Essence'],
    'Breach Hives': ['Breach'],
}

# Items named in each tab besides the waystone and the tablet slots (those two are read from their cells).
# (role, market name, kind, label, count, wording in the tab[, mods the item must have])
#   kind c: currency-exchange item   u: unique   g: gem   b: base type   grp: a kind of loot, not one item
OMENS3 = [('in', 'Omen of Chaotic Rarity', 'c', '', 1, 'Omen of Chaotic Rarity'),
          ('in', 'Omen of Chaotic Quantity', 'c', '', 1, 'Rarity, Quantity and Effectiveness'),
          ('in', 'Omen of Chaotic Effectiveness', 'c', '', 1, 'Rarity, Quantity and Effectiveness'),
          ('in', 'Exalted Orb', 'c', '', 1, 'Exalt afterwards')]
ITEMS = {
    'Deli Unstable Breach': [
        ('in', 'Omen of Chaotic Quantity', 'c', '', 1, 'Omen of Chaotic Quantity'),
        ('in', 'Omen of Chaotic Monsters', 'c', '', 1, 'Chaotic Monsters'),
        ('out', 'Hiveblood', 'grp', '', 1, 'hiveblood'),
        ('out', 'Catalysts', 'grp', '', 1, 'Catalysts'),
        ('out', 'Rings', 'grp', '', 1, 'rings/amulets'),
        ('out', 'Amulets', 'grp', '', 1, 'rings/amulets')],
    'Lineage/Fragment Farming': [
        ('in', 'Tablets with Waystone chance', 'grp', '', 3, '3x Any Tablets with % Inc chance of Waystones',
         ['% Inc chance of Waystones']),
        ('in', 'Visions of Paradise', 'u', 'Tablet', 1, '1x Visions of Pradise'),
        ('in', 'Mastered Domain', 'u', 'Tablet', 1, '1x Forest Mastered Domain', ['Forest biome']),
        ('out', "Rakiata's Flow", 'g', '', 1, 'Rakiatas Flow'),
        ('out', "Garukhan's Resolve", 'g', '', 1, 'Garukhans Resolve'),
        ('out', 'Other Lineage gems', 'grp', '', 1, 'Lineage Gems'),
        ('out', 'Citadel boss fragments', 'grp', '', 1, 'Citadel Bosses drop valuable fragments')],
    'Deli Abyss': OMENS3 + [('out', 'Omens', 'grp', '', 1, 'high tier omens')],
    'Jackpot Ritual': [('in', 'Head of the King', 'c', '', 1, 'Use the Head of the King')],
    'Ritual': [('in', 'Head of the King', 'c', '', 1, 'Use the Head of the King'),
               ('out', 'Omens', 'grp', '', 1, 'Expensive Omens')],
    'Tablet Farming': [('out', 'Tablets', 'grp', '', 1, 'tablets are expensive')],
    'Aldurs Saga Hunting': [('in', 'Expedition Logbook', 'c', '', 1, 'use a logbook'),
                            ('out', "Aldur's Saga", 'c', 'Expedition', 1, 'Aldurs Saga')],
    'Big Boom Expedition': [('out', 'Expedition Logbook', 'c', '', 1, 'Logbooks drop')],
    'Grand Expedition Farming': [],
    'Hilda Logbook Farming': [('out', 'Expedition Logbook', 'c', '', 1, 'Logbooks drop'),
                              ('out', "Aldur's Saga", 'c', 'Expedition', 1, 'hunt for Aldurs Saga')],
    'Unstable Hiveblood farming': [
        ('in', 'Wombgifts (ring/amulet)', 'grp', '', 1, 'ring/amulet wombgifts'),
        ('out', 'Hiveblood', 'grp', '', 1, 'farming hiveblood'),
        ('out', 'Absent Amulet', 'b', 'Amulet', 1, 'Absent amulets'),
        ('out', 'Catalysts', 'grp', '', 1, 'catalysts'),
        ('out', 'Rings', 'grp', '', 1, 'rings'),
        ('out', 'Amulets', 'grp', '', 1, 'amulets')],
    'Deli Boss Rush': [],
    '200% Doryani Deli Farming': OMENS3,
    'Deli Splinter Farming': OMENS3 + [('out', 'Simulacrum Splinter', 'c', '', 1, 'sim splinters')],
    'Liquid Emotion Farming': [('out', 'Liquid emotions', 'grp', '', 1, 'liquid emotions'),
                               ('out', 'Vaal currency', 'grp', '', 1, 'vaal currency')],
    'Abyssal Boxes': [('out', 'Chest currency', 'grp', '', 1, 'chest currency')],
    'Simulacrum': [('in', 'Simulacrum', 'c', '', 1, 'with the fragment'),
                   ('out', 'Voices', 'u', 'Jewel', 1, 'unique Voices jewel')],
    'Atziri Boss Rush': [],
}
# Lineage/Fragment Farming has its own layout: three map setups, the atlas master of each shown only as a
# screenshot of the master's passives (Doryani's Science above Citadel Maps, Jado's Spycraft above the other two).
LINEAGE_SETUPS = [('B15', 'B17', 'B19', 'Doryani'), ('B28', 'B30', 'B32', 'Jado'), ('B41', 'B43', 'B45', 'Jado')]

# Mod-specific inputs (rolled tablets and waystones) have their own price: each gets an official trade
# search, written to data/farmqueries.json. The site's worker runs them through the hour and serves the
# results at data/farmprices.json. Plain items keep their poe.ninja price.
# Sheet wording -> trade stat (ids from data/trade.json, the trade site's own list). Two ids = either one.
# A number in the wording is the minimum; "reduced" is a negative "increased" value.
STAT_RULES = [
    (r'(\d+)(?:-\d+)?\+? additional rare monsters', ('explicit.stat_3762913035', 'explicit.stat_1653625239'), 'rare'),
    (r'(\d+) additional map modifiers', ('explicit.stat_588512487',), 'mods'),
    (r'rerolling favours (\d+) additional times', ('explicit.stat_120737942',), 'reroll'),
    (r'deferring favours costs % reduced tribute', ('explicit.stat_1345835998',), 'defer-cheap'),
    (r'rerolling favours costs % reduced tribute', ('explicit.stat_2282052746',), 'reroll-cheap'),
    (r'favours to be omens', ('explicit.stat_4219853180',), 'omens'),
    (r'(?:quantity|chance) of waystones', ('explicit.stat_2777224821',), 'waystones'),
    (r'monsters have %\s*inc effectiveness|^monster effectiveness', ('explicit.stat_2065500219',), 'effect'),
    (r'rarity of items found', ('explicit.stat_2306002879',), 'rarity'),
    (r'fracturing mirrors', ('explicit.stat_551040294',), 'frac'),
    (r'sim splinters', ('explicit.stat_3836551197',), 'splinters'),
    (r'monster rarity', ('explicit.stat_4142653832',), 'monster-rarity'),
    (r'% inc rare monsters', ('explicit.stat_3793155082',), 'rares'),
    (r'quantity of hiveblood', ('explicit.stat_2778285247',), 'hiveblood'),
    (r'(?:(\d+)%\+?\s*)?more likely to spawn unique bosses', ('explicit.stat_3962960008',), 'bosses'),
    (r'mirror shards', ('explicit.stat_900933517',), 'shards'),
    (r'abyss pits .*twice as likely', ('explicit.stat_4256531808',), 'pits'),
    (r'chance to contain additional abysses', ('explicit.stat_2890355696',), 'four-abysses'),
    (r'(?:(\d+)%\+?\s*)?chance to add a vaal beacon unique monster', ('explicit.stat_3937291366',), 'vaal-unique'),
    (r'additional crystal from vaal beacons', ('explicit.stat_1940774881',), 'crystal'),
    (r'\bforest\b', ('explicit.stat_864099561',), 'forest'),
]
OPTIONAL = re.compile(r'^\((?:and if you can|try)\)', re.I)   # nice to have, not part of the search

URL_RE = re.compile(r'https?://\S+')
problems = []


def warn(msg):
    problems.append(msg)


def norm(s):
    return re.sub(r'[^a-z0-9]', '', (s or '').lower())


def fix(text):
    for pat, rep, fl in TYPOS:
        text = re.sub(pat, rep, text, flags=fl)
    return text


def line(v):
    """One tidy line of cell text, typos fixed."""
    return fix(' '.join(str(v).split())) if v is not None else None


def cell(ws, ref):
    v = ws[ref].value
    return line(v) if isinstance(v, str) and v.strip() else None


def link(ws, ref):
    h = ws[ref].hyperlink
    return h.target if h is not None and h.target else None


def fetch(url):
    req = urllib.request.Request(url, headers={'User-Agent': UA})
    with urllib.request.urlopen(req, timeout=180) as r:
        return r.read()


def tab_list():
    """Tab name -> gid (for links to each strategy's page) and the document title, from the public html view."""
    try:
        h = fetch(BASE + '/htmlview').decode('utf-8', 'replace')
    except Exception as e:  # links then point at the sheet's first tab
        warn('tab list not loaded: %s' % e)
        return {}, None
    t = re.search(r'<title>(.*?)</title>', h, re.S)
    title = re.sub(r'\s*-\s*Google (Drive|Sheets)\s*$', '', html.unescape(t.group(1)).strip()) if t else None
    gids = {norm(n.replace('\\/', '/')): g for n, g in re.findall(r'\{name: "(.*?)", pageUrl: ".*?gid=(\d+)', h)}
    return gids, title


def trade_lists():
    """Tablet base types and the base of each unique, from the official trade site's own lists (data/trade.json)."""
    t = json.loads((ROOT / 'data' / 'trade.json').read_text(encoding='utf-8'))
    return [x for x in t['bases'].get('Maps', []) if x.endswith(' Tablet')], t.get('uniques', {})


def stars(text, full):
    return text.count(full) if text else None


def note_and_links(text):
    """A note with its pasted links taken out. A short lead-in right before a link ("Catalyst Tree Link:")
    becomes the button label. Returns (note or None, [(label, url)])."""
    parts, urls = URL_RE.split(text), URL_RE.findall(text)
    links = []
    for i, u in enumerate(urls):
        before = re.sub(r'[\s:\-]+$', '', parts[i])
        cut = max(before.rfind('. '), before.rfind(' - '))
        lead = before[cut + 2:] if cut >= 0 else before
        lead = lead.strip(' -')
        if lead and len(lead) <= 60 and not lead.startswith('('):
            parts[i] = before[:len(before) - len(lead)] if cut >= 0 else ''
        else:
            lead = ''
        label = next((v for k, v in LINK_LABELS.items() if k in u), None) or lead or 'Link'
        if all(u != x[1] for x in links):
            links.append((label, u))
    rest = ' '.join(''.join(parts).split()).strip(' -:')
    rest = re.sub(r'^\((.*)\)$', r'\1', rest)
    if len(rest.split()) < 3:
        return None, links
    if not re.search(r'[.!?)"]$', rest):
        rest += '.'
    return rest, links


def bullet(v):
    return re.sub(r'^•\s*', '', v).strip()


def tablets_of(ws, refs):
    """Tablet slots, same tablet and mods merged with a count. Returns [] when the tab says none are used."""
    out = []
    for ref in refs:
        raw = ws[ref].value
        if not isinstance(raw, str) or not raw.strip():
            continue
        lines = [line(x) for x in raw.split('\n') if x.strip()]
        name, mods = lines[0], [m for m in lines[1:] if not CUT_LINE.search(m)]
        if len(mods) < len(lines) - 1:
            warn('%s %s: left out a cut-off tablet line' % (ws.title, ref))
        if name.rstrip('.').lower() in ('n/a', 'not needed', 'none'):
            continue
        if out and out[-1]['n'] == name and out[-1]['mods'] == mods:
            out[-1]['x'] += 1
        else:
            out.append({'n': name, 'x': 1, 'mods': mods})
    return out


def slug(s):
    return re.sub(r'[^a-z0-9]+', '-', s.lower().replace("'", '')).strip('-')


class Queries:
    """The trade searches for mod-specific inputs: one per distinct search, under a short readable key."""
    def __init__(self):
        self.by_key = {}

    def add(self, key, label, query):
        body = {'query': query, 'sort': {'price': 'asc'}}
        have = self.by_key.get(key)
        if have and have['query'] != body:
            warn('two different searches under one key: ' + key)
        self.by_key.setdefault(key, {'key': key, 'label': label, 'query': body})
        return key


def mod_stats(lines, quiet=False):
    """The trade stats an item must have, from its mod lines. "(and if you can)" and "(try)" lines are extras.
    quiet: a unique's lines often just describe its own roll, so no stat there is no problem."""
    found, codes, used = [], [], []
    for ln in lines:
        if OPTIONAL.match(ln):
            continue
        low = re.sub(r'^\(need\)\s*', '', ln.lower())
        hit = False
        for pat, ids, code in STAT_RULES:
            m = re.search(pat, low)
            if not m or any(ids == f[0] for f in found):
                continue
            num = next((g for g in m.groups() if g), None)
            found.append((ids, {'max': -1} if 'reduced' in pat else ({'min': int(num)} if num else None)))
            codes.append(code + (num or ''))
            hit = True
        if hit:
            used.append(re.sub(r'^\(need\)\s*', '', ln).rstrip(' ,'))
        elif not quiet:
            warn('no trade stat for "%s"' % ln)
    return found, codes, used


def stat_groups(found):
    ands, groups = [], []
    for ids, val in found:
        flt = [dict({'id': i}, **({'value': val} if val else {})) for i in ids]
        if len(flt) == 1:
            ands += flt
        else:   # either wording counts
            groups.append({'type': 'count', 'filters': flt, 'value': {'min': 1}})
    return ([{'type': 'and', 'filters': ands}] if ands else []) + groups


def tablet_query(Q, uniq, base, lines):
    """A trade search for a tablet with the mods the sheet asks for; None when it asks for none."""
    found, codes, used = mod_stats(lines, quiet=bool(uniq))
    if not found:
        return None
    q = {'status': {'option': 'online'}}
    if uniq:
        q['name'] = uniq
        if base:
            q['type'] = base
    elif base:
        q['type'] = base
        q['filters'] = {'type_filters': {'filters': {'rarity': {'option': 'nonunique'}}}}
    else:   # any tablet
        q['filters'] = {'type_filters': {'filters': {'category': {'option': 'map.tablet'}, 'rarity': {'option': 'nonunique'}}}}
    q['stats'] = stat_groups(found)
    name = uniq or base or 'Any tablet'
    return Q.add(slug(name) + '-' + '-'.join(codes), name + ': ' + '; '.join(used), q)


def waystone_input(Q, text):
    """The waystone a strategy runs, with a trade search for the mods, corruption and drop chance it asks for."""
    m = re.search(r'\bT(\d{1,2})\b', text or '')
    if not m:
        return None
    name = 'Waystone (Tier %s)' % m.group(1)
    t = text.lower()
    q = {'status': {'option': 'online'}, 'type': name}
    stats, filters, mapf, codes, words = [], {}, {}, ['t' + m.group(1)], []
    mods = re.search(r'(\d)\s*mod', t)
    if mods:
        n = int(mods.group(1))
        exact = n < 6   # "HAS TO BE 5 MODS"; 6 is the most a waystone has
        stats.append({'id': 'pseudo.pseudo_number_of_affix_mods', 'value': {'min': n, 'max': n} if exact else {'min': n}})
        codes.append('%dmod' % n)
        words.append(('exactly %d mods' if exact else '%d mods') % n)
    corrupted = True if 'corrupted' in t else (False if mods and 'craft' in t else None)   # crafting needs it clean
    if corrupted is not None:
        filters['misc_filters'] = {'filters': {'corrupted': {'option': 'true' if corrupted else 'false'}}}
        codes.append('corrupted' if corrupted else 'clean')
        words.append('corrupted' if corrupted else 'not corrupted')
    drop = re.search(r'(\d+)%\+?\s*waystone', t)
    if drop:
        mapf['map_bonus'] = {'min': int(drop.group(1))}
        codes.append('drop' + drop.group(1))
        words.append(drop.group(1) + '%+ Waystone Drop Chance')
    if 'monster effectiveness' in t:
        mapf['map_magic_monsters'] = {'min': 1}
        codes.append('effect')
        words.append('Monster Effectiveness')
    if 'rarity of items' in t:
        mapf['map_iir'] = {'min': 1}
        codes.append('rarity')
        words.append('Item Rarity')
    if mapf:
        filters['map_filters'] = {'filters': mapf}
    if stats:
        q['stats'] = [{'type': 'and', 'filters': stats}]
    if filters:
        q['filters'] = filters
    key = Q.add('waystone-' + '-'.join(codes), name + (': ' + ', '.join(words) if words else ''), q)
    return {'n': name, 'k': 'c', 's': 'Waystone', 'm': [text], 'tk': key}


def parse_standard(ws, name, author):
    f = {}
    f['updated'] = re.sub(r'^Updated\s+', '', cell(ws, 'Z2') or '') or None
    tier = cell(ws, 'L6')
    f['_tier_on_page'] = tier
    master = cell(ws, 'B29')
    if master:
        f['master'] = master
    rows = {}
    merged = list(ws.merged_cells.ranges)
    for c in ws[28]:   # row 28 holds the box titles, row 29 their values
        if not isinstance(c.value, str):
            continue
        rng = next((m for m in merged if m.min_row == 28 and m.min_col == c.column), None)
        lo, hi = (rng.min_col, rng.max_col) if rng else (c.column, c.column)
        val = next((x.value for x in ws[29] if lo <= x.column <= hi and isinstance(x.value, str) and x.value.strip()), None)
        rows[c.value.strip().lower()] = val
    if rows.get('difficulty'):
        f['d'] = stars(rows['difficulty'], '★')
    if rows.get('investment'):
        f['i'] = stars(rows['investment'], '✦')
    f['waystone'] = cell(ws, 'AV11')
    f['tablets'] = tablets_of(ws, ('AU19', 'AU22', 'AU24', 'AU26'))
    if rows.get('rarity'):
        f['rarity'] = line(rows['rarity'])
    tip = rows.get('recommendation') or rows.get('psa')
    if tip:
        f['tip'] = line(tip)
    links, notes, trade, extra = [], [], [], []
    vid = cell(ws, 'D9')
    if vid and link(ws, 'D9'):
        who = re.sub(r'^Vid by\s+', '', vid)
        links.append({'k': 'video', 't': 'Video by ' + (author if re.match(r'^Me\b', who) else who), 'u': link(ws, 'D9')})
    guide = cell(ws, 'D12')
    if guide and link(ws, 'D12'):
        if name in WRONG_GUIDE:
            warn('%s: guide link left out (points to %s)' % (name, link(ws, 'D12')))
        else:
            links.append({'k': 'guide', 't': guide, 'u': link(ws, 'D12')})
    if link(ws, 'D16'):
        links.append({'k': 'tree', 't': 'Atlas tree', 'u': link(ws, 'D16')})
    for r in range(11, 28):
        for col in ('M', 'W', 'AG'):
            ref = '%s%d' % (col, r)
            raw = ws[ref].value
            if not isinstance(raw, str) or not raw.strip():
                continue
            v, u = line(raw), link(ws, ref)
            if u and 'pathofexile.com/trade2' in u:
                trade.append({'k': 'trade', 't': re.sub(r'\s*(Trade\s+)?Link\b', '', bullet(v)).strip(), 'u': u})
            elif u and not v.startswith('•'):
                extra.append({'k': 'link', 't': next((l for k, l in LINK_LABELS.items() if k in u), v), 'u': u})
            else:
                n, found = note_and_links(bullet(v))
                if n:
                    notes.append(n)
                if u and all(u != uu for _, uu in found):
                    found.append((next((l for k, l in LINK_LABELS.items() if k in u), 'Link'), u))
                extra += [{'k': 'link', 't': lab, 'u': uu} for lab, uu in found]
    f['notes'] = notes
    f['links'] = links + trade + extra
    return f


def parse_lineage(ws, name, author):
    f = {}
    f['updated'] = re.sub(r'^Updated\s+', '', cell(ws, 'Z2') or '') or None
    title = ws['I4'].value or ''
    f['_tier_on_page'] = title.strip().split('\n')[-1].strip()
    f['setups'] = []
    for map_ref, way_ref, tab_ref, master in LINEAGE_SETUPS:
        f['setups'].append({'map': cell(ws, map_ref), 'master': master, 'waystone': cell(ws, way_ref),
                            'tablets': re.sub(r'^Tablets:\s*', '', cell(ws, tab_ref) or '')})
    f['master'] = ', '.join(dict.fromkeys(s['master'] for s in f['setups']))
    f['waystone'] = f['setups'][0]['waystone']
    f['tablets'] = []
    links, notes, extra = [], [], []
    for r in range(9, 46):
        ref = 'M%d' % r
        raw = ws[ref].value
        if not isinstance(raw, str) or not raw.strip():
            continue
        v, u = line(raw), link(ws, ref)
        if not v.startswith('•'):
            if u and 'maxroll.gg/poe2/atlas-tree' in u:
                links.append({'k': 'tree', 't': 'Atlas tree', 'u': u})
            continue
        m = re.match(r'^(?:Great )?video by (\S+)', bullet(v), re.I)
        if m and u:
            links.insert(0, {'k': 'video', 't': 'Video by ' + m.group(1), 'u': u})
            continue
        n, found = note_and_links(bullet(v))
        if n:
            notes.append(n)
        extra += [{'k': 'link', 't': lab, 'u': uu} for lab, uu in found]
    f['notes'] = notes
    f['links'] = links + extra
    return f


def main():
    if len(sys.argv) > 1:
        data = Path(sys.argv[1]).read_bytes()
    else:
        print('downloading the sheet…')
        data = fetch(BASE + '/export?format=xlsx')
    wb = openpyxl.load_workbook(io.BytesIO(data))
    gids, doc_title = tab_list()
    by_norm = {norm(ws.title): ws for ws in wb.worksheets}   # the xlsx drops "/" from tab names
    tab_name = {norm(k): k for k in list(MECH) + list(ITEMS)}
    bases, uniques = trade_lists()
    Q = Queries()

    tl = by_norm['tierlist']
    author = cell(tl, 'AA5') or 'BawLoch'
    labels = [(c.row, c.column, c.value.strip().split()[0]) for row in tl.iter_rows() for c in row
              if isinstance(c.value, str) and re.fullmatch(r'\s*[SABCDF] Tier\s*', c.value)]
    entries = []
    for row in tl.iter_rows(min_row=min(r for r, _, _ in labels)):
        for c in row:
            if not isinstance(c.value, str) or not c.value.strip() or re.fullmatch(r'\s*[SABCDF] Tier\s*', c.value):
                continue
            same = [l for l in labels if l[0] <= c.row]
            top = max(l[0] for l in same)
            tier = max((l for l in same if l[0] == top and l[1] <= c.column), key=lambda l: l[1])[2]
            target = c.hyperlink.location if c.hyperlink is not None and c.hyperlink.location else None
            entries.append((line(c.value), tier, target.split('!')[0].strip("'") if target else None))

    farms, leagues = [], collections.Counter()
    for name, tier, target in entries:
        f = {'id': slug(name), 'name': name, 'tier': tier}
        ws = by_norm.get(norm(target)) if target else None
        key = tab_name.get(norm(ws.title if ws else name))
        f['mech'] = MECH.get(key, [])
        if not ws:
            f.update({'notes': [], 'in': [], 'out': [], 'links': []})
            farms.append(f)
            continue
        if not key:
            warn('%s: new tab, add it to MECH and ITEMS' % ws.title)
        standard = str(ws['L6'].value or '').strip().endswith('Tier')
        body = (parse_standard if standard else parse_lineage)(ws, key or ws.title, author)
        on_page = body.pop('_tier_on_page', '') or ''
        if not on_page.startswith(tier + ' '):
            warn('%s: tier on its page is "%s", tier list says %s' % (name, on_page, tier))
        f.update(body)
        for u in (l['u'] for l in f['links'] if l['k'] == 'trade'):
            m = re.search(r'/trade2/search/poe2/([^/]+)/', u)
            if m:
                leagues[urllib.parse.unquote(m.group(1))] += 1
        gid = gids.get(norm(ws.title)) or gids.get(norm(key))
        if gid:
            f['links'].append({'k': 'sheet', 't': 'Sheet page', 'u': BASE + '/edit?gid=' + gid})

        # what it spends: the waystone, the tablet slots, then the items named in the tab
        text = '\n'.join(str(c.value) for r in ws.iter_rows() for c in r if c.value is not None).lower()
        ins, outs, tabs = [], [], []
        w = waystone_input(Q, f.get('waystone'))
        if w:
            ins.append(w)
        for t in f['tablets']:   # one input per slot kind: a tablet with other mods is another item
            base = next((b for b in bases if t['n'].endswith(b)), None)
            if not base:
                warn('%s: tablet "%s" has no known base' % (name, t['n']))
                continue
            uniq = t['n'][:-len(base)].strip()
            item = {'n': uniq or base, 'k': 'u' if uniq else 'b', 's': 'Tablet', 'q': t['x'], 'm': t['mods']}
            tk = tablet_query(Q, uniq, base, t['mods'])
            if tk:
                item['tk'] = tk
            have = next((x for x in tabs if x['n'] == item['n'] and x.get('tk') == tk and x['m'] == item['m']), None)
            if have:
                have['q'] += t['x']
            else:
                tabs.append(item)
        for role, n, kind, sub, q, wording, *mods in ITEMS.get(key, []):
            if wording.lower() not in text:
                warn('%s: "%s" not found in the tab any more (for %s)' % (name, wording, n))
            dest = ins if role == 'in' else outs
            if any(x['n'] == n for x in dest):
                continue
            item = {'n': n, 'k': kind}
            if sub:
                item['s'] = sub
            if q != 1:
                item['q'] = q
            if mods:
                item['m'] = mods[0]
                base = (uniques.get(n) or [None])[0] if kind == 'u' else None
                tk = tablet_query(Q, n if kind == 'u' else None, base, mods[0])
                if tk:
                    item['tk'] = tk
            dest.append(item)
        ins += tabs   # waystone, then the crafting items, then the tablets
        for x in ins:
            if x.get('q') == 1:
                del x['q']
        f['in'], f['out'] = ins, outs
        farms.append(f)

    order = {t: i for i, t in enumerate('SABCDF')}
    farms.sort(key=lambda f: order.get(f['tier'], 9))   # stable: keeps the sheet's own order inside a tier
    c4 = tl['C4'].value or ''
    ver = re.search(r'\b(\d+\.\d+)\b', c4)
    upd = re.sub(r'^Updated\s+', '', cell(tl, 'R2') or '')
    try:
        upd_iso = dt.datetime.strptime(upd, '%d %B %Y').date().isoformat()
    except ValueError:
        upd_iso = None
    alinks = []
    for ref in ('AJ6', 'AO6', 'AT6', 'T5'):
        u = link(tl, ref)
        if u:
            alinks.append({'t': next((l for d, l in AUTHOR_LINKS if d in u), 'Link'), 'u': u})
    source = {
        'title': doc_title or line(c4),
        'version': ver.group(1) if ver else None,
        'league': leagues.most_common(1)[0][0] if leagues else None,   # not in any cell: read from its trade links
        'author': author,
        'links': alinks,
        'updated': upd or None,
        'updated_iso': upd_iso,
        'url': BASE + '/edit' + ('?gid=' + gids['tierlist'] if 'tierlist' in gids else ''),
        'fetched': dt.date.today().isoformat(),
    }
    body = ',\n'.join(json.dumps(f, ensure_ascii=False, separators=(',', ':')) for f in farms)
    with open(OUT, 'w', encoding='utf-8', newline='\n') as fh:
        fh.write('{"source":' + json.dumps(source, ensure_ascii=False, separators=(',', ':')) +
                 ',\n"farms":[\n' + body + '\n]}\n')
    queries = sorted(Q.by_key.values(), key=lambda x: x['key'])
    with open(QUERIES, 'w', encoding='utf-8', newline='\n') as fh:
        fh.write('{"updated":"%s","queries":[\n' % source['fetched'] +
                 ',\n'.join(json.dumps(x, ensure_ascii=False, separators=(',', ':')) for x in queries) + '\n]}\n')
    print('%d farms (%s), league %s, updated %s -> %s' % (
        len(farms), ' '.join('%s:%d' % kv for kv in collections.Counter(f['tier'] for f in farms).items()),
        source['league'], source['updated'], OUT.relative_to(ROOT)))
    print('%d trade searches -> %s' % (len(queries), QUERIES.relative_to(ROOT)))
    for p in problems:
        print('  note:', p)


if __name__ == '__main__':
    main()
